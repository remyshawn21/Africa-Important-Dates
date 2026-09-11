"""
NSE1 Consolidator
=================
Merges all country Excel files into the master "All Important Dates" sheet
of the workbook read by update_dashboard.py.

Usage:
    python consolidate.py

Setup:
    1. Put all country Excel files in the 'country_files' folder
    2. Run this script — it creates/updates 'Africa Important Dates.xlsx',
       writing into its "All Important Dates" sheet. If a "Big Ticket
       Activities" sheet already exists in that file, it is left untouched
       (Big Ticket entries are curated by hand, not merged from country files).

Each country file must have these columns:
    Country | Date | Event | Status | Description | Category of Event | Requires PR

After running this, run update_dashboard.py to rebuild index.html.
"""

COUNTRY_FOLDER = "country_files"
MASTER_FILE    = "Africa Important Dates.xlsx"
MAIN_SHEET     = "All Important Dates"
BIGTICKET_SHEET = "Big Ticket Activities"

import sys, os, glob
from datetime import datetime

def check_dependencies():
    missing = []
    try: import pandas
    except ImportError: missing.append("pandas")
    try: import openpyxl
    except ImportError: missing.append("openpyxl")
    if missing:
        print(f"\n❌  Missing: pip install {' '.join(missing)}\n")
        sys.exit(1)

check_dependencies()
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REQUIRED_COLS = {'Country', 'Date', 'Event', 'Status', 'Description',
                  'Category of Event', 'Requires PR'}
COL_ORDER = ['Country', 'Date', 'Event', 'Status', 'Description',
             'Category of Event', 'Requires PR']

def load_country_files():
    if not os.path.exists(COUNTRY_FOLDER):
        print(f"\n❌  Folder '{COUNTRY_FOLDER}' not found.")
        print(f"    Create it and put country Excel files inside.\n")
        sys.exit(1)

    files = glob.glob(os.path.join(COUNTRY_FOLDER, "*.xlsx"))
    if not files:
        print(f"\n❌  No .xlsx files found in '{COUNTRY_FOLDER}'\n")
        sys.exit(1)

    frames, errors = [], []
    for f in files:
        name = os.path.basename(f)
        try:
            df = pd.read_excel(f)
            missing = REQUIRED_COLS - set(df.columns)
            if missing:
                errors.append(f"  ⚠️  {name} — missing columns: {sorted(missing)}")
                continue
            df['_source'] = name
            frames.append(df)
            print(f"  ✅  {name} — {len(df)} rows")
        except Exception as e:
            errors.append(f"  ❌  {name} — error: {e}")

    if errors:
        print("\nWarnings:")
        for e in errors: print(e)

    if not frames:
        print("\n❌  No valid files loaded.\n")
        sys.exit(1)

    return pd.concat(frames, ignore_index=True)

def validate_and_clean(df):
    today = pd.Timestamp.now().normalize()

    df['_raw_date']         = df['Date']  # keep original text (e.g. "TBC", "September")
    df['Date']              = pd.to_datetime(df['Date'], errors='coerce')
    df['Description']       = df['Description'].fillna('').astype(str).str.strip()
    df['Status']            = df['Status'].fillna('').astype(str).str.strip()
    df['Event']             = df['Event'].fillna('').astype(str).str.strip()
    df['Country']           = df['Country'].fillna('').astype(str).str.strip()
    df['Category of Event'] = df['Category of Event'].fillna('').astype(str).str.strip()
    df['Requires PR']       = df['Requires PR'].fillna('No').astype(str).str.strip()

    df = df[df['Event'] != '']

    n_unparsed = df['Date'].isna().sum()
    if n_unparsed:
        print(f"\n  ℹ️   {n_unparsed} row(s) have a date that needs update_dashboard.py's"
              f" own date parser (not pandas) to interpret, e.g. 'TBC' or a month"
              f" name — their original text is kept as-is.")

    issues = []

    # Rule 1: Executed but date is in future → revert (only checkable for parsed dates)
    future_executed = (df['Status'].str.lower() == 'executed') & (df['Date'] > today)
    if future_executed.any():
        issues.append(f"  ⚠️  {future_executed.sum()} event(s) marked Executed but date hasn't passed → reverted to 'Yet to Happen'")
        df.loc[future_executed, 'Status'] = 'Yet to Happen'

    # Rule 2: Hide descriptions for future events (only checkable for parsed dates)
    future_with_desc = (df['Date'] > today) & (df['Description'] != '')
    if future_with_desc.any():
        issues.append(f"  ℹ️   {future_with_desc.sum()} future event(s) had descriptions → hidden until date passes")
        df.loc[future_with_desc, 'Description'] = ''

    if issues:
        print("\nValidation:")
        for i in issues: print(i)

    return df

def style_sheet(ws, n_cols):
    header_fill = PatternFill("solid", start_color="1B3F8B", end_color="1B3F8B")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    thin   = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    for cell in ws[1]:
        cell.fill = header_fill; cell.font = header_font
        cell.alignment = center; cell.border = border

    exec_fill = PatternFill("solid", start_color="E8F5EE", end_color="E8F5EE")
    post_fill = PatternFill("solid", start_color="FFF3E8", end_color="FFF3E8")
    yet_fill  = PatternFill("solid", start_color="E6ECF8", end_color="E6ECF8")

    # Column D = Status (index 4)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        status = str(row[3].value or '').lower()
        fill = exec_fill if 'executed' in status else post_fill if 'postponed' in status else yet_fill
        for cell in row:
            cell.fill = fill; cell.border = border
            cell.font = Font(name="Arial", size=9)
            cell.alignment = left if cell.column in [3, 5] else center

    widths = {'A': 16, 'B': 14, 'C': 38, 'D': 16, 'E': 50, 'F': 22, 'G': 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

def save_master(df):
    out = df[COL_ORDER + ['_raw_date']].copy()
    # Sort by parsed date (rows with an unparseable date sort to the end),
    # then by Country.
    out = out.sort_values(['Date', 'Country'], na_position='last').reset_index(drop=True)
    out['Date'] = out.apply(
        lambda r: r['Date'].strftime('%d-%b-%Y') if pd.notna(r['Date']) else str(r['_raw_date']),
        axis=1
    )
    out = out[COL_ORDER]

    # If the master file already exists, preserve any other sheets in it
    # (most importantly "Big Ticket Activities", which isn't built from
    # country files and shouldn't be overwritten here).
    if os.path.exists(MASTER_FILE):
        wb = load_workbook(MASTER_FILE)
        if MAIN_SHEET in wb.sheetnames:
            del wb[MAIN_SHEET]
        ws = wb.create_sheet(MAIN_SHEET, 0)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = MAIN_SHEET
        if BIGTICKET_SHEET not in wb.sheetnames:
            bt = wb.create_sheet(BIGTICKET_SHEET)
            bt.append(COL_ORDER)
            style_sheet(bt, len(COL_ORDER))

    ws.append(COL_ORDER)
    for row in out.itertuples(index=False):
        ws.append(list(row))

    style_sheet(ws, len(COL_ORDER))
    wb.save(MASTER_FILE)

def main():
    print("\n── NSE1 Consolidator ───────────────────────────────")
    print(f"  Scanning: {COUNTRY_FOLDER}/\n")

    df = load_country_files()
    df = validate_and_clean(df)
    save_master(df)

    total     = len(df)
    countries = df['Country'].nunique()
    executed  = df['Status'].str.lower().str.contains('executed').sum()
    yet       = df['Status'].str.lower().str.contains('yet').sum()
    with_desc = (df['Description'] != '').sum()
    with_pr   = (df['Requires PR'].str.lower() == 'yes').sum()

    print(f"\n  Countries: {countries}  |  Events: {total}")
    print(f"  Executed: {executed}  |  Upcoming: {yet}  |  With reports: {with_desc}  |  Requires PR: {with_pr}")
    print(f"  Saved:    {MASTER_FILE}  (sheet: '{MAIN_SHEET}')")
    print("────────────────────────────────────────────────────")
    print("  ✅  Master file ready! Now run update_dashboard.py")
    print("────────────────────────────────────────────────────\n")

if __name__ == "__main__":
    main()
